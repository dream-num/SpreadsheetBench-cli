#!/usr/bin/env python3
import argparse
import json
import re
from pathlib import Path

from openpyxl import load_workbook


def split_answer_position(text):
    parts = []
    current = []
    quote = False
    for ch in text:
        if ch == "'":
            quote = not quote
        if ch == "," and not quote:
            part = "".join(current).strip()
            if part:
                parts.append(part)
            current = []
        else:
            current.append(ch)
    part = "".join(current).strip()
    if part:
        parts.append(part)
    return parts


def parse_ref(ref):
    ref = ref.strip()
    if "!" in ref:
        sheet, cell_range = ref.rsplit("!", 1)
        sheet = sheet.strip().strip("'")
    else:
        sheet = None
        cell_range = ref
    cell_range = cell_range.strip().strip("'")
    return sheet, cell_range


def value_key(v):
    return v


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("task_id")
    parser.add_argument("--run-id", default="codex-gpt-5-5-verified400-all-20260526-231834")
    parser.add_argument("--dataset", default="data/spreadsheetbench_verified_400")
    parser.add_argument("--limit", type=int, default=80)
    args = parser.parse_args()

    dataset = Path(args.dataset)
    dataset_json = json.loads((dataset / "dataset.json").read_text())
    meta = next(row for row in dataset_json if row["id"] == args.task_id)

    output = dataset / "outputs" / "univer_agent_gpt-5.5" / f"1_{args.task_id}_output.xlsx"
    golden = dataset / "spreadsheet" / args.task_id / f"1_{args.task_id}_golden.xlsx"
    if not golden.exists():
        golden = dataset / "spreadsheet" / args.task_id / "golden.xlsx"

    print(f"task_id: {args.task_id}")
    print(f"instruction_type: {meta.get('instruction_type')}")
    print(f"answer_position: {meta.get('answer_position')}")
    print(f"output: {output}")
    print(f"golden: {golden}")
    print()

    out_wb = load_workbook(output, data_only=True)
    gold_wb = load_workbook(golden, data_only=True)
    total = 0
    shown = 0
    for ref in split_answer_position(meta["answer_position"]):
        sheet, cell_range = parse_ref(ref)
        if sheet is None:
            sheet = gold_wb.sheetnames[0]
        if sheet not in gold_wb.sheetnames or sheet not in out_wb.sheetnames:
            print(f"[sheet-missing] {ref}: output_has={sheet in out_wb.sheetnames} golden_has={sheet in gold_wb.sheetnames}")
            continue
        gs = gold_wb[sheet]
        os = out_wb[sheet]
        print(f"[range] {sheet}!{cell_range}")
        try:
            g_cells = gs[cell_range]
            o_cells = os[cell_range]
        except Exception as exc:
            print(f"[range-error] {exc}")
            continue
        if not isinstance(g_cells, tuple):
            g_cells = ((g_cells,),)
            o_cells = ((o_cells,),)
        for g_row, o_row in zip(g_cells, o_cells):
            for g_cell, o_cell in zip(g_row, o_row):
                if value_key(g_cell.value) != value_key(o_cell.value):
                    total += 1
                    if shown < args.limit:
                        print(f"{g_cell.coordinate}: output={o_cell.value!r} golden={g_cell.value!r}")
                        shown += 1
    print()
    print(f"diff_count_shown_or_more: {total}; shown: {shown}; limit: {args.limit}")


if __name__ == "__main__":
    main()
