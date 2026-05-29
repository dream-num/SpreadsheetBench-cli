from pathlib import Path
from typing import Dict, Iterable, Optional

from openpyxl import load_workbook

from .paths import task_id_text


AGENT_PROMPT_TEMPLATE = """You are a spreadsheet expert helping a user complete a workbook editing request inside a Docker container.
Use `benchmarking-univer-cli` before touching any workbook: `skill: benchmarking-univer-cli`.

The user provided one or more workbook cases. For each case, edit the workbook according to the request and create the required output.xlsx file.

This prompt is only the dynamic task envelope. Follow `benchmarking-univer-cli` for SaC-only mutation, benchmark evaluator semantics, Univer Facade pitfalls, and plan/assertion/verify/export discipline.

The request contains these types of information:
- instruction: The user's workbook editing request.
- spreadsheet_path: The prepared SaC workspace and managed artifact paths you need to manipulate.
- spreadsheet_content: A first-rows preview of the first input spreadsheet file.
- instruction_type: Cell-Level Manipulation or Sheet-Level Manipulation.
- answer_position: The evaluator-facing target/check range for the final workbook state.
- output_path: The required modified spreadsheet files.

Request id: {task_id}

### instruction
{instruction}

### spreadsheet_path
{case_lines}

### spreadsheet_content
{spreadsheet_content}

### instruction_type
{instruction_type}

### answer_position
{answer_position}

### output_path
{output_lines}

Benchmark task envelope:
- Only use files under /task.
- Treat the listed /task/cases/case_N/sac paths as prepared SaC workspaces for solving.
- The listed managed artifacts are the workbook artifacts controlled by those SaC workspaces.
- Create the required /task/outputs/case_N/output.xlsx file for every case.
- Keep temporary scripts and intermediates under /task/work.
- Solve every case independently.
- Follow `benchmarking-univer-cli` for SaC-only mutation, answer_position interpretation, readonly probes, and final export gates.
"""


def build_spreadsheet_content(input_file: Path, max_rows: int = 5) -> str:
    workbook = load_workbook(input_file, data_only=False, read_only=True)
    sections = []
    for sheet in workbook.worksheets:
        rows = []
        for row in sheet.iter_rows(max_row=max_rows, values_only=True):
            rows.append("\t".join("" if value is None else str(value) for value in row))
        sections.append(
            "Sheet Name: " + sheet.title + "\n"
            + "\n".join(rows)
            + "\n"
            + "-" * 50
        )
    workbook.close()
    return "\n".join(sections)


def build_agent_prompt(
    task: Dict,
    cases: Optional[Iterable[int]] = None,
    spreadsheet_content: str = "",
) -> str:
    case_list = list(cases or [1])
    case_lines = "\n".join(
        "\n".join(
            [
                f"- /task/cases/case_{case_index}/sac: prepared SaC workspace for case {case_index}.",
                f"  Managed artifact: /task/cases/case_{case_index}/sac/artifacts/sac.univer",
            ]
        )
        for case_index in case_list
    )
    output_lines = "\n".join(
        f"- /task/outputs/case_{case_index}/output.xlsx"
        for case_index in case_list
    )
    return AGENT_PROMPT_TEMPLATE.format(
        task_id=task_id_text(task),
        instruction=task.get("instruction", ""),
        instruction_type=task.get("instruction_type", ""),
        answer_position=task.get("answer_position", ""),
        spreadsheet_content=spreadsheet_content,
        case_lines=case_lines,
        output_lines=output_lines,
    )
