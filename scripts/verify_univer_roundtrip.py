#!/usr/bin/env python3
"""Verify Univer CLI import/export roundtrip for SpreadsheetBench xlsx files."""

from __future__ import annotations

import argparse
import csv
import datetime as dt
import gc
import json
import shutil
import subprocess
import sys
import tempfile
import time
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import openpyxl


@dataclass(frozen=True)
class CaseFile:
    task_id: str
    kind: str
    source_path: Path


def discover_tasks(dataset_dir: Path, task_ids: set[str] | None) -> list[tuple[str, list[CaseFile]]]:
    spreadsheet_dir = dataset_dir / "spreadsheet"
    tasks: list[tuple[str, list[CaseFile]]] = []
    for task_dir in sorted(p for p in spreadsheet_dir.iterdir() if p.is_dir()):
        task_id = task_dir.name
        if task_ids and task_id not in task_ids:
            continue

        init_files = sorted(task_dir.glob(f"*_{task_id}_init.xlsx"))
        golden_files = sorted(task_dir.glob(f"*_{task_id}_golden.xlsx"))
        if not init_files and (task_dir / "initial.xlsx").is_file():
            init_files = [task_dir / "initial.xlsx"]
        if not golden_files and (task_dir / "golden.xlsx").is_file():
            golden_files = [task_dir / "golden.xlsx"]

        cases = [CaseFile(task_id, "input", path) for path in init_files]
        cases.extend(CaseFile(task_id, "output", path) for path in golden_files)
        if cases:
            tasks.append((task_id, cases))
    return tasks


def select_cases(
    dataset_dir: Path,
    task_ids: set[str] | None,
    batch_index: int | None,
    batch_size: int | None,
    limit_files: int | None,
) -> list[CaseFile]:
    tasks = discover_tasks(dataset_dir, task_ids)
    if batch_index is not None:
        if batch_size is None:
            raise ValueError("--batch-index requires --batch-size")
        start = (batch_index - 1) * batch_size
        end = start + batch_size
        tasks = tasks[start:end]

    cases = [case for _, task_cases in tasks for case in task_cases]
    if limit_files is not None:
        cases = cases[:limit_files]
    return cases


def run_cmd(argv: list[str], timeout: int, log_path: Path) -> tuple[bool, str]:
    started = time.monotonic()
    try:
        result = subprocess.run(
            argv,
            capture_output=True,
            text=True,
            timeout=timeout,
        )
        elapsed = time.monotonic() - started
        log_path.write_text(
            json.dumps(
                {
                    "argv": argv,
                    "returncode": result.returncode,
                    "elapsed_seconds": round(elapsed, 3),
                    "stdout": result.stdout,
                    "stderr": result.stderr,
                },
                ensure_ascii=False,
                indent=2,
            )
            + "\n",
            encoding="utf-8",
        )
        if result.returncode != 0:
            return False, result.stderr.strip() or result.stdout.strip()
        return True, ""
    except subprocess.TimeoutExpired as exc:
        elapsed = time.monotonic() - started
        log_path.write_text(
            json.dumps(
                {
                    "argv": argv,
                    "timeout": timeout,
                    "elapsed_seconds": round(elapsed, 3),
                    "stdout": exc.stdout,
                    "stderr": exc.stderr,
                },
                ensure_ascii=False,
                indent=2,
            )
            + "\n",
            encoding="utf-8",
        )
        return False, f"timeout after {timeout}s"


def comparable_value(value: Any) -> Any:
    if value == "":
        return None
    if isinstance(value, float):
        return round(value, 10)
    if isinstance(value, dt.datetime):
        return value.isoformat(sep=" ")
    if isinstance(value, dt.date):
        return value.isoformat()
    if isinstance(value, dt.time):
        return value.isoformat()
    return value


def worksheet_values(ws: Any) -> dict[tuple[int, int], Any]:
    values: dict[tuple[int, int], Any] = {}
    for (row, col), cell in ws._cells.items():
        value = comparable_value(cell.value)
        if value is not None:
            values[(row, col)] = value
    return values


def compare_loaded_workbooks(src_path: Path, exported_path: Path) -> tuple[bool, str]:
    src_wb = None
    exported_wb = None
    try:
        src_wb = openpyxl.load_workbook(src_path, data_only=True)
        exported_wb = openpyxl.load_workbook(exported_path, data_only=True)
        if src_wb.sheetnames != exported_wb.sheetnames:
            return False, f"sheetnames differ: src={src_wb.sheetnames}, exported={exported_wb.sheetnames}"

        for sheet_name in src_wb.sheetnames:
            src_ws = src_wb[sheet_name]
            exported_ws = exported_wb[sheet_name]
            src_keys: set[tuple[int, int]] = set()
            for (row, col), src_cell in sorted(src_ws._cells.items()):
                src_value = comparable_value(src_cell.value)
                if src_value is None:
                    continue
                src_keys.add((row, col))
                exported_value = comparable_value(exported_ws.cell(row=row, column=col).value)
                if src_value != exported_value:
                    coord = openpyxl.utils.get_column_letter(col) + str(row)
                    return (
                        False,
                        f"value diff at {sheet_name}!{coord}: src={src_value!r}, exported={exported_value!r}",
                    )
            for (row, col), exported_cell in sorted(exported_ws._cells.items()):
                if (row, col) in src_keys:
                    continue
                exported_value = comparable_value(exported_cell.value)
                if exported_value is not None:
                    coord = openpyxl.utils.get_column_letter(col) + str(row)
                    return (
                        False,
                        f"value diff at {sheet_name}!{coord}: src=None, exported={exported_value!r}",
                    )
        return True, ""
    finally:
        if src_wb is not None:
            src_wb.close()
        if exported_wb is not None:
            exported_wb.close()
        gc.collect()


def check_formula_text(src_path: Path, exported_path: Path) -> tuple[bool, str]:
    src_wb = None
    exported_wb = None
    try:
        src_wb = openpyxl.load_workbook(src_path, data_only=False)
        exported_wb = openpyxl.load_workbook(exported_path, data_only=False)
        if src_wb.sheetnames != exported_wb.sheetnames:
            return False, "skipped because sheetnames differ"

        for sheet_name in src_wb.sheetnames:
            src_ws = src_wb[sheet_name]
            exported_ws = exported_wb[sheet_name]
            src_formula_keys: set[tuple[int, int]] = set()
            for (row, col), src_cell in sorted(src_ws._cells.items()):
                src_formula = src_cell.value
                if not (isinstance(src_formula, str) and src_formula.startswith("=")):
                    continue
                src_formula_keys.add((row, col))
                exported_formula = exported_ws.cell(row=row, column=col).value
                if src_formula != exported_formula:
                    coord = openpyxl.utils.get_column_letter(col) + str(row)
                    return (
                        False,
                        f"formula diff at {sheet_name}!{coord}: "
                        f"src={src_formula!r}, exported={exported_formula!r}",
                    )
            for (row, col), exported_cell in sorted(exported_ws._cells.items()):
                if (row, col) in src_formula_keys:
                    continue
                exported_formula = exported_cell.value
                if isinstance(exported_formula, str) and exported_formula.startswith("="):
                    coord = openpyxl.utils.get_column_letter(col) + str(row)
                    return (
                        False,
                        f"formula diff at {sheet_name}!{coord}: src=None, exported={exported_formula!r}",
                    )
        return True, ""
    finally:
        if src_wb is not None:
            src_wb.close()
        if exported_wb is not None:
            exported_wb.close()
        gc.collect()


def verify_one(
    case: CaseFile,
    out_dir: Path,
    temp_root: Path,
    timeout: int,
    keep_artifacts: bool,
    keep_failed_artifacts: bool,
) -> dict[str, Any]:
    safe_name = f"{case.task_id}-{case.kind}-{case.source_path.stem}"
    logs_dir = out_dir / "logs" / safe_name
    logs_dir.mkdir(parents=True, exist_ok=True)

    artifact_dir = out_dir / "artifacts" / safe_name
    temp_root.mkdir(parents=True, exist_ok=True)

    row: dict[str, Any] = {
        "task_id": case.task_id,
        "kind": case.kind,
        "source_path": str(case.source_path),
        "artifact_dir": "",
        "status": "ok",
        "import_ok": False,
        "export_ok": False,
        "openpyxl_ok": False,
        "values_ok": False,
        "formulas_ok": None,
        "message": "",
    }

    try:
        with tempfile.TemporaryDirectory(prefix=f"{safe_name}-", dir=temp_root) as case_dir_name:
            case_dir = Path(case_dir_name)
            univer_path = case_dir / f"{safe_name}.univer"
            exported_path = case_dir / f"{safe_name}.roundtrip.xlsx"

            import_ok, import_msg = run_cmd(
                ["univer", "import", str(case.source_path), str(univer_path), "--json"],
                timeout,
                logs_dir / "import.json",
            )
            row["import_ok"] = import_ok
            if not import_ok:
                row["status"] = "import_error"
                row["message"] = import_msg
            else:
                export_ok, export_msg = run_cmd(
                    ["univer", "export", str(univer_path), str(exported_path), "--json"],
                    timeout,
                    logs_dir / "export.json",
                )
                row["export_ok"] = export_ok
                if not export_ok:
                    row["status"] = "export_error"
                    row["message"] = export_msg
                else:
                    try:
                        values_ok, values_msg = compare_loaded_workbooks(case.source_path, exported_path)
                        row["openpyxl_ok"] = True
                        row["values_ok"] = values_ok
                        if not values_ok:
                            row["status"] = "value_mismatch"
                            row["message"] = values_msg
                    except Exception as exc:
                        row["status"] = "openpyxl_error"
                        row["message"] = repr(exc)

                    if row["status"] != "openpyxl_error":
                        try:
                            formulas_ok, formulas_msg = check_formula_text(case.source_path, exported_path)
                            row["formulas_ok"] = formulas_ok
                            if not formulas_ok and row["status"] == "ok":
                                row["status"] = "formula_mismatch"
                                row["message"] = formulas_msg
                        except Exception as exc:
                            row["formulas_ok"] = False
                            if row["status"] == "ok":
                                row["status"] = "formula_check_error"
                                row["message"] = repr(exc)

            retain = keep_artifacts or (keep_failed_artifacts and row["status"] != "ok")
            if retain:
                if artifact_dir.exists():
                    shutil.rmtree(artifact_dir, ignore_errors=True)
                shutil.copytree(case_dir, artifact_dir)
                row["artifact_dir"] = str(artifact_dir)
    except Exception as exc:
        row["status"] = "script_error"
        row["message"] = repr(exc)
    finally:
        gc.collect()

    return row


def write_reports(out_dir: Path, rows: list[dict[str, Any]]) -> None:
    summary: dict[str, Any] = {
        "total": len(rows),
        "by_status": {},
        "by_kind": {},
        "failed": [row for row in rows if row["status"] != "ok"],
    }
    for row in rows:
        summary["by_status"][row["status"]] = summary["by_status"].get(row["status"], 0) + 1
        kind_counts = summary["by_kind"].setdefault(row["kind"], {})
        kind_counts[row["status"]] = kind_counts.get(row["status"], 0) + 1

    (out_dir / "summary.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    (out_dir / "results.json").write_text(
        json.dumps(rows, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    with (out_dir / "results.csv").open("w", encoding="utf-8", newline="") as fh:
        writer = csv.DictWriter(fh, fieldnames=list(rows[0].keys()) if rows else ["status"])
        writer.writeheader()
        writer.writerows(rows)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--dataset-dir",
        default="data/spreadsheetbench_verified_400",
        type=Path,
        help="SpreadsheetBench dataset directory.",
    )
    parser.add_argument(
        "--out-dir",
        type=Path,
        default=None,
        help="Output directory. Defaults to .runs/cli-roundtrip-ver400-YYYYMMDD-HHMMSS.",
    )
    parser.add_argument("--task-id", action="append", help="Task id to include; may repeat.")
    parser.add_argument("--limit", type=int, default=None, help="Limit number of xlsx files after task batching.")
    parser.add_argument("--batch-index", type=int, default=None, help="1-based task batch index.")
    parser.add_argument("--batch-size", type=int, default=None, help="Number of tasks in each batch.")
    parser.add_argument(
        "--temp-root",
        type=Path,
        default=Path("/private/tmp/univer-roundtrip"),
        help="Temporary artifact root; each case is deleted after verification.",
    )
    parser.add_argument("--timeout", type=int, default=180, help="Per import/export timeout seconds.")
    parser.add_argument("--keep-artifacts", action="store_true", help="Keep .univer/.xlsx for every case.")
    parser.add_argument("--keep-failed-artifacts", action="store_true", help="Keep .univer/.xlsx only for failed cases.")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    out_dir = args.out_dir
    if out_dir is None:
        run_id = "cli-roundtrip-ver400-" + dt.datetime.now().strftime("%Y%m%d-%H%M%S")
        out_dir = Path(".runs") / run_id
    out_dir.mkdir(parents=True, exist_ok=True)

    task_ids = set(args.task_id) if args.task_id else None
    cases = select_cases(args.dataset_dir, task_ids, args.batch_index, args.batch_size, args.limit)

    rows: list[dict[str, Any]] = []
    batch_label = ""
    if args.batch_index is not None:
        batch_label = f" batch={args.batch_index} batch_size={args.batch_size}"
    print(f"dataset={args.dataset_dir} files={len(cases)}{batch_label} out_dir={out_dir}", flush=True)
    for index, case in enumerate(cases, 1):
        print(f"[{index}/{len(cases)}] {case.task_id} {case.kind} {case.source_path}", flush=True)
        row = verify_one(
            case,
            out_dir,
            args.temp_root,
            args.timeout,
            args.keep_artifacts,
            args.keep_failed_artifacts,
        )
        rows.append(row)
        write_reports(out_dir, rows)
        if row["status"] != "ok":
            print(f"  -> {row['status']}: {row['message']}", flush=True)

    write_reports(out_dir, rows)
    failed = [row for row in rows if row["status"] != "ok"]
    print(f"done total={len(rows)} failed={len(failed)} summary={out_dir / 'summary.json'}")
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
